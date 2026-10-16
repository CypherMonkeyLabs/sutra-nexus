"""Pastel ~15 slide deck via python-pptx. Output: build/downloads/sutra-nexus-slide-deck.pptx. Native charts from data. No external assets."""
import os, json, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
from nexus_lib import *
C=json.loads(open(f'{ROOT}/js/content.js',encoding='utf8').read().split('=',1)[1].rstrip().rstrip(';'))
M=load('maheshvara');A=load('ashtadhyayi');K=load('shiva_sutras_kashmir');S={x['num']:x for x in A['sutras']}
INK=RGBColor(0x1c,0x1b,0x3a);MUTE=RGBColor(0x47,0x4b,0x66);C1=RGBColor(0x5b,0x21,0xb6);C2=RGBColor(0x0e,0x74,0x90);C3=RGBColor(0xbe,0x18,0x5d)
BGS=[RGBColor(0xee,0xfb,0xf5),RGBColor(0xf1,0xea,0xff),RGBColor(0xff,0xee,0xf5),RGBColor(0xe8,0xf2,0xff),RGBColor(0xff,0xf6,0xdc)]
FOOT='Sūtra Nexus · Ancient rules. Living systems. · © 2026 CypherMonkeyLabs · Independent, free, learning-only student project. Trademarks and sources belong to their owners. Concerns: https://cyphermonkeylabs.com/#contact · Not sold.'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6];n=[0]
def tb(sl,x,y,w,h,text,size=18,bold=False,color=INK,font='Calibri',align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    s=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=s.text_frame;tf.word_wrap=True;tf.vertical_anchor=anchor
    lines=text if isinstance(text,list) else [text]
    for i,l in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;r=p.add_run();r.text=l;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=color;r.font.name=font;p.space_after=Pt(6)
    return s
def slide(title,sub=None):
    sl=prs.slides.add_slide(blank);n[0]+=1
    bg=sl.background.fill;bg.solid();bg.fore_color.rgb=BGS[n[0]%len(BGS)]
    bar=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,Inches(.18));bar.fill.solid();bar.fill.fore_color.rgb=[C1,C2,C3][n[0]%3];bar.line.fill.background()
    tb(sl,.6,.35,12,.9,title,34,True,C1)
    if sub: tb(sl,.6,1.15,12,.5,sub,16,False,MUTE)
    f=tb(sl,.6,6.85,10.9,.55,FOOT,9,False,MUTE);tb(sl,12,6.85,.8,.4,str(n[0]),11,True,C1,align=PP_ALIGN.RIGHT)
    return sl
def table(sl,x,y,w,rows,cols,size=12,colw=None,rh=.35):
    t=sl.shapes.add_table(len(rows)+1,len(cols),Inches(x),Inches(y),Inches(w),Inches(rh*(len(rows)+1))).table
    for j,c in enumerate(cols):
        cell=t.cell(0,j);cell.text=c;cell.fill.solid();cell.fill.fore_color.rgb=C1
        for p in cell.text_frame.paragraphs:
            for r in p.runs:r.font.size=Pt(size);r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            cell=t.cell(i+1,j);cell.text=str(v);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(255,255,255) if i%2==0 else RGBColor(0xf6,0xf2,0xff)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:r.font.size=Pt(size);r.font.color.rgb=INK;r.font.name='Noto Sans Devanagari' if any('\u0900'<=ch<='\u097f' for ch in str(v)) else 'Calibri'
    for r_ in t.rows: r_.height=Inches(rh)
    if colw:
        for j,wd in enumerate(colw):t.columns[j].width=Inches(wd)
    return t
# 1 title
sl=slide('Sūtra Nexus','Ancient rules. Living systems.')
tb(sl,.6,2.2,12,3,['Sanskrit sound and grammar, read from open data.','A free learning project by a lifelong student and AI enthusiast.','Māheśvara sūtras · Aṣṭādhyāyī · Kashmir Śiva Sūtras'],24,False,INK)
tb(sl,.6,5.4,12,.6,'ॐ   अ इ उ ण्   वृद्धिरादैच्   हल्',32,True,C2,'Noto Sans Devanagari')
# 2 what
sl=slide('What this is','A study companion that runs offline in a browser')
tb(sl,.6,1.9,12,4.5,['• 10 rooms: Start Here, Māheśvara sūtras, Aṣṭādhyāyī, Śiva Sūtras, Sound lab, Practice, Ecosystem, Resources, Downloads, About.','• A pratyāhāra builder computed from the 14 lines and checked against known sets.','• A browser for 3,983 Aṣṭādhyāyī sūtras with anuvṛtti and charts.','• Flashcards with a spaced ladder, quizzes, a daily challenge and a guide bot.','• No internet needed. Progress stays in your browser.'],20)
# 3 data
sl=slide('Data and sources','Only what the open files contain')
table(sl,.6,1.9,12.1,[['Aṣṭādhyāyī','3,983 sūtras','ashtadhyayi-com/data','Free to use with credit'],['Māheśvara sūtras','14 lines','ashtadhyayi-com/data','Free to use with credit'],['Śiva Sūtras','75 sūtras (20+10+45)','GRETIL, ed. Madhusudan Kaul Shastri 1925, input O. Hellwig','Credit the inputter; one catalogue lists CC BY-NC-SA 4.0']],['Text','Count','Source','Licence note'],13,[2.4,2.6,4.2,2.9])
tb(sl,.6,4.2,12,2,['Transliteration (IAST) was converted by a script and has not been checked by a person.','English meanings are not in the data: meaning: UNKNOWN in source. Edition and translator for any gloss: UNKNOWN.'],16,False,MUTE)
# 4 loop
sl=slide('The learning loop','Sound, structure, meaning, use')
for i,(a,b) in enumerate([('Sound','Say it aloud'),('Structure','See how the rule is built'),('Meaning','Read it, labelled as interpretation'),('Use','Try one drill')]):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.6+i*3.1),Inches(2.4),Inches(2.8),Inches(2),);sh.fill.solid();sh.fill.fore_color.rgb=[C1,C2,C3,RGBColor(0x04,0x78,0x57)][i];sh.line.fill.background()
    tf=sh.text_frame;tf.word_wrap=True;p=tf.paragraphs[0];p.alignment=PP_ALIGN.CENTER;r=p.add_run();r.text=a;r.font.size=Pt(26);r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    p2=tf.add_paragraph();p2.alignment=PP_ALIGN.CENTER;r=p2.add_run();r.text=b;r.font.size=Pt(15);r.font.color.rgb=RGBColor(255,255,255)
# 5 14 lines
sl=slide('The 14 Māheśvara lines','Each ends with a marker that is not counted')
table(sl,.6,1.55,12.1,[[l['n'],l['devanagari'],l['iast'],' '.join(l['sounds']),l['marker']] for l in M['sutras']],['Line','Devanāgarī','IAST','Sounds','Marker'],10,[.9,3.2,3.2,3.6,1.2],rh=.3)
# 6 how
sl=slide('How a pratyāhāra works','Sūtra 1.1.71 ādirantyena sahetā')
tb(sl,.6,1.9,12,4.5,['A first sound plus a final marker names itself and every sound between them.','Example: ac starts at a in line 1 and stops at the marker c of line 4.','Result: a i u ṛ ḷ e o ai au (9 sounds). The marker is not counted.','Another example: hal runs from h in line 5 to the marker l of line 14: 33 consonants.','Checked in the app against ac, ik, yaṇ, śar, ak, ec, hal and al.'],20)
# 7 table
sl=slide('Common pratyāhāras','Computed from the 14 lines')
rows=[]
for nm in ['ak','ac','ec','ik','yaṇ','jhal','hal','al','śar']:
    p=[x for x in C['prat'] if x['name']==nm][0];rows.append([nm,len(p['sounds']),' '.join(p['sounds'])[:70]])
table(sl,.6,1.8,12.1,rows,['Name','Count','Sounds'],13,[1.6,1.4,9.1])
# 8 sthana
sl=slide('Where sounds are made','Sthāna: textbook summary')
table(sl,.6,1.8,12.1,[[a,b] for a,b,_ in C['sthana']],['Sthāna','Sounds'],13,[5.1,7])
# 9 sandhi
sl=slide('Four sandhi flows','Rules from the data; glosses are interpretations')
table(sl,.6,1.9,12.1,[[s[0],s[1],S[s[0]]['devanagari']+' ('+S[s[0]]['iast']+')',s[2]+' → '+s[3],s[4]] for s in C['sandhi']],['Sūtra','Name','Rule','Example','Change'],12,[1.2,1.3,4,3.6,2])
# 10 chart adhyaya
sl=slide('Sūtras per adhyāya','Computed from the data: '+f"{C['stats']['total']:,} in all")
cd=CategoryChartData();cd.categories=[f'Adhyāya {k}' for k in C['stats']['adhyaya']];cd.add_series('Sūtras',list(C['stats']['adhyaya'].values()))
ch=sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.8),Inches(1.7),Inches(11.5),Inches(4.9),cd).chart;ch.has_legend=False;ch.plots[0].series[0].format.fill.solid();ch.plots[0].series[0].format.fill.fore_color.rgb=C1;ch.plots[0].has_data_labels=True
# 11 types
sl=slide('Sūtra types','S, P and V are expanded in the data meta; AD and AT are not')
cd=CategoryChartData();T=C['stats']['types'];cd.categories=['V vidhi','S saṃjñā','AT','AD','P paribhāṣā'];cd.add_series('Count',[T['V'],T['S'],T['AT'],T['AD'],T['P']])
ch=sl.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT,Inches(.8),Inches(1.7),Inches(6.5),Inches(4.9),cd).chart;ch.has_legend=True;ch.legend.position=XL_LEGEND_POSITION.RIGHT;ch.legend.include_in_layout=False
tb(sl,7.6,2.2,5.2,4,[f"V: {T['V']:,}",f"S: {T['S']}",f"AT: {T['AT']}",f"AD: {T['AD']}",f"P: {T['P']}"],20)
# 12 anuvritti
sl=slide('Anuvṛtti','Words carried forward so short rules stay short')
rows=[[n,S[n]['devanagari'],' · '.join(S[n]['anuvritti']) or 'none'] for n in ['1.1.3','1.1.4','1.1.5','1.1.60','6.1.87','6.1.101']]
table(sl,.6,1.9,12.1,rows,['Sūtra','Text','Carried forward (from the data)'],14,[1.6,4.4,6.1])
# 13 shiva
sl=slide('Śiva Sūtras','75 sūtras in this file: 20 + 10 + 45')
rows=[[x['num'],x['devanagari'],x['iast']] for x in K['sutras'][:6]]
table(sl,.6,1.9,12.1,rows,['No.','Devanāgarī','IAST'],14,[1.2,5,5.9])
tb(sl,.6,4.8,12,1.8,['The file’s meta note says the common Kṣemarāja recension counts 22 + 10 + 45. Which sūtras differ: UNKNOWN beyond the source.','No translation is included; meaning: UNKNOWN in source.'],15,False,MUTE)
# 14 honesty
sl=slide('Honesty rules','What is and is not in this project')
tb(sl,.6,1.9,12,4.6,['• Every canon line carries a citation: work, locus and source.','• Interpretations are labelled as bridge ideas.','• No invented sūtras, verses, meanings, editions, translators or dates.','• The Bhagavad Gita is not in the data and is not included.','• Gaps: English meanings, human-checked transliteration, the 75 versus 77 difference.'],20)
# 15 about
sl=slide('About and links','Independent, free, learning-only')
tb(sl,.6,1.9,12,4.5,['CypherMonkey Labs: https://cyphermonkeylabs.com','GitHub organization: https://github.com/CypherMonkeyLabs','GitHub: https://github.com/cipher0x9','X: https://x.com/Cypher0x9','LinkedIn: https://www.linkedin.com/in/luckylakshmanucwebex','Hugging Face: https://huggingface.co/Cypher0x9'],18)
out=f'{ROOT}/downloads/sutra-nexus-slide-deck.pptx';prs.save(out);print(out,os.path.getsize(out),len(prs.slides._sldIdLst))
